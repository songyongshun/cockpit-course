"""探针：智能座舱课件的逐张幻灯片概览（栏目横幅/分隔页/重复文本）。"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation

SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2])
MSO_GROUP = 6


def iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def clean(line: str) -> str:
    return line.replace("\u000b", " ").replace("\u3000", " ").replace("\xa0", " ").strip()


with OUT.open("w", encoding="utf-8") as fh:
    for path in sorted(SRC.rglob("*.pptx")):
        prs = Presentation(str(path))
        fh.write(f"\n### {path.parent.name} / {path.name}  ({path.stat().st_size / 1e6:.1f} MB)\n")
        counter: Counter[str] = Counter()
        for no, slide in enumerate(prs.slides, start=1):
            texts, pics, tables = [], 0, 0
            for shape in iter_shapes(slide.shapes):
                if getattr(shape, "has_table", False) and shape.has_table:
                    tables += 1
                if getattr(shape, "has_text_frame", False):
                    text = (shape.text_frame.text or "").strip()
                    if text:
                        texts.append(text)
                if shape.shape_type == 13:
                    pics += 1
            flat = " / ".join(t.replace("\n", " ⏎ ") for t in texts)[:150]
            for text in texts:
                counter[clean(text)] += 1
            fh.write(f"S{no:02d} chars={len(''.join(texts)):5d} pics={pics} tbl={tables} "
                     f"shp={len(texts):2d} | {flat}\n")
        fh.write("  -- 重复文本块(>=3): --\n")
        for text, n in counter.most_common():
            if n >= 3 and text:
                fh.write(f"    x{n:3d} | {text[:70]!r}\n")
print("done")